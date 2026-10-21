from __future__ import annotations

import os
from pathlib import Path
from math import sin, cos, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\LoRA")
OUT_DIR = ROOT / "output"
ASSET_DIR = ROOT / "report_assets"
OUT_DIR.mkdir(exist_ok=True)
ASSET_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "番茄棚关键参数与采摘机器人场景调研.docx"

FONT_REG = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"

GREEN = "356859"
DARK = "1F2933"
MID = "607D72"
PALE = "EEF5F1"
PALE_BLUE = "EDF3F7"
LIGHT = "F7F9F8"
GRAY = "6B7280"
BORDER = "D9D9D9"
RED = "D9534F"
ORANGE = "E69F45"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def draw_arrow(draw, start, end, fill, width=5, head=14):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    angle = __import__("math").atan2(y2 - y1, x2 - x1)
    for a in (angle + 2.55, angle - 2.55):
        draw.line([(x2, y2), (x2 + head * cos(a), y2 + head * sin(a))], fill=fill, width=width)


def make_height_figure(path: Path):
    im = Image.new("RGB", (1600, 1050), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 48), "植株高度与结果高度口径", font=font(54, True), fill="#111111")
    d.text((72, 118), "建议以果实中心距地面的垂直距离作为结果高度", font=font(28), fill="#4B5563")

    ground_y = 900
    base_x = 830
    top_y = 225
    d.line([(130, ground_y), (1490, ground_y)], fill="#7A6750", width=8)
    d.text((1320, ground_y + 24), "地面", font=font(28), fill="#5B4636")

    # stem and leaves
    d.line([(base_x, ground_y), (base_x, top_y)], fill="#2D6A4F", width=17)
    for i, y in enumerate(range(top_y + 70, ground_y - 30, 85)):
        direction = -1 if i % 2 == 0 else 1
        d.line([(base_x, y), (base_x + direction * 160, y - 42)], fill="#3D7C5B", width=10)
        cx, cy = base_x + direction * 190, y - 55
        d.ellipse((cx - 72, cy - 34, cx + 72, cy + 34), fill="#79A96B", outline="#46784F", width=4)

    # fruits at several representative heights
    fruit_centers = [(690, 705), (955, 650), (700, 555), (960, 500), (700, 395), (950, 335)]
    fruit_colors = ["#C83E36", "#E15943", "#F0A33A", "#D74439", "#E8763A", "#D34A3E"]
    for (x, y), c in zip(fruit_centers, fruit_colors):
        d.line([(base_x, y - 15), (x, y)], fill="#3D7C5B", width=6)
        d.ellipse((x - 38, y - 35, x + 38, y + 35), fill=c, outline="#9C2B26", width=3)
        d.polygon([(x - 15, y - 32), (x, y - 47), (x + 15, y - 32), (x, y - 25)], fill="#39764B")

    # canopy height arrow
    xh = 1290
    draw_arrow(d, (xh, ground_y), (xh, top_y), "#374151", width=5, head=18)
    draw_arrow(d, (xh, top_y), (xh, ground_y), "#374151", width=5, head=18)
    d.rounded_rectangle((1330, 480, 1540, 645), radius=18, fill="#F3F4F6", outline="#D1D5DB", width=3)
    d.text((1360, 505), "植株高度", font=font(31, True), fill="#111111")
    d.text((1360, 555), "1.5–1.7 m", font=font(34), fill="#356859")
    d.text((1360, 602), "默认 1.6 m", font=font(28), fill="#4B5563")

    # fruit zone bracket
    y_low, y_high = 665, 340
    xb = 300
    d.line([(xb, y_high), (xb, y_low)], fill="#B83A32", width=8)
    d.line([(xb - 22, y_high), (xb + 22, y_high)], fill="#B83A32", width=8)
    d.line([(xb - 22, y_low), (xb + 22, y_low)], fill="#B83A32", width=8)
    d.rounded_rectangle((60, 405, 260, 605), radius=18, fill="#FFF3F0", outline="#E6B7AE", width=3)
    d.text((88, 430), "成熟果区", font=font(31, True), fill="#8F2C26")
    d.text((82, 485), "0.55–1.40 m", font=font(31), fill="#B83A32")
    d.text((98, 540), "默认重点区", font=font(25), fill="#4B5563")
    d.text((86, 578), "0.60–1.30 m", font=font(27), fill="#4B5563")

    d.text((70, 960), "说明：植株高度指仿真中地面至冠层顶端的可见高度；结果高度指果实中心距地面高度。", font=font(26), fill="#4B5563")
    im.save(path, quality=95)


def make_layout_figure(path: Path):
    im = Image.new("RGB", (1600, 950), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 45), "番茄棚行列布局与机器人通道", font=font(52, True), fill="#111111")
    d.text((72, 112), "俯视示意  数值为场景生成推荐值", font=font(28), fill="#4B5563")

    # crop rows
    row_xs = [260, 550, 1050, 1340]
    for x in row_xs:
        d.rounded_rectangle((x - 65, 210, x + 65, 820), radius=26, fill="#DCEED9", outline="#7AA36E", width=4)
        for y in range(260, 800, 92):
            d.ellipse((x - 34, y - 34, x + 34, y + 34), fill="#6EA661", outline="#3E7447", width=3)
            d.ellipse((x - 10, y - 8, x + 10, y + 12), fill="#D54B3F")

    # narrow and wide spacing annotations
    d.line([(260, 855), (550, 855)], fill="#374151", width=4)
    draw_arrow(d, (300, 855), (510, 855), "#374151", width=4, head=14)
    draw_arrow(d, (510, 855), (300, 855), "#374151", width=4, head=14)
    d.text((302, 866), "小行距 0.4–0.6 m", font=font(24), fill="#374151")

    # aisle
    aisle_left, aisle_right = 660, 940
    d.rounded_rectangle((aisle_left, 205, aisle_right, 825), radius=28, fill="#EDF3F7", outline="#93A7B5", width=4)
    d.line([(800, 230), (800, 800)], fill="#9AAAB5", width=4)
    for y in range(260, 780, 70):
        d.line([(786, y), (814, y)], fill="#9AAAB5", width=4)
    # robot
    d.rounded_rectangle((705, 465, 895, 625), radius=24, fill="#536B7A", outline="#2F4451", width=5)
    d.ellipse((680, 495, 730, 555), fill="#283640")
    d.ellipse((870, 495, 920, 555), fill="#283640")
    d.ellipse((680, 560, 730, 620), fill="#283640")
    d.ellipse((870, 560, 920, 620), fill="#283640")
    d.text((745, 520), "底盘", font=font(34, True), fill="white")
    d.polygon([(800, 450), (783, 430), (817, 430)], fill="#E69F45")
    d.text((746, 655), "预设采摘点", font=font(25), fill="#2F4451")
    d.text((700, 690), "每 0.8 m 设一个起始点", font=font(24), fill="#2F4451")

    # aisle width arrow
    d.line([(aisle_left, 170), (aisle_right, 170)], fill="#356859", width=5)
    draw_arrow(d, (aisle_left + 30, 170), (aisle_right - 30, 170), "#356859", width=5, head=15)
    draw_arrow(d, (aisle_right - 30, 170), (aisle_left + 30, 170), "#356859", width=5, head=15)
    d.text((680, 125), "净通道 1.0–1.2 m", font=font(30, True), fill="#356859")

    # plant spacing label
    d.line([(1410, 330), (1410, 422)], fill="#374151", width=4)
    draw_arrow(d, (1410, 340), (1410, 412), "#374151", width=4, head=12)
    draw_arrow(d, (1410, 412), (1410, 340), "#374151", width=4, head=12)
    d.text((1435, 340), "株距", font=font(26, True), fill="#374151")
    d.text((1435, 380), "默认 0.40 m", font=font(24), fill="#374151")

    d.text((70, 900), "工程关系：通道净宽应不小于底盘宽度加两侧安全余量；最终值需用实际底盘外廓复核。", font=font(25), fill="#4B5563")
    im.save(path, quality=95)


def make_range_figure(path: Path):
    im = Image.new("RGB", (1600, 950), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 45), "核心参数的文献范围与推荐默认值", font=font(50, True), fill="#111111")
    d.text((72, 108), "彩色线段表示文献或标准覆盖范围  圆点表示建议默认值", font=font(27), fill="#4B5563")

    rows = [
        ("株距", "m", 0.20, 0.60, 0.30, 0.50, 0.40),
        ("植株高度", "m", 0.0, 2.0, 1.50, 1.70, 1.60),
        ("成熟果高度", "m", 0.0, 1.6, 0.55, 1.40, 0.95),
        ("单株留果量", "个", 0, 45, 16, 40, 24),
        ("果实横径", "mm", 50, 110, 62.2, 98.1, 78.4),
    ]
    y0 = 235
    for i, (name, unit, amin, amax, low, high, default) in enumerate(rows):
        y = y0 + i * 135
        d.text((90, y - 22), name, font=font(31, True), fill="#1F2933")
        d.text((90, y + 24), f"范围 {low:g}–{high:g} {unit}", font=font(24), fill="#4B5563")
        x1, x2 = 560, 1450
        d.line([(x1, y + 12), (x2, y + 12)], fill="#D1D5DB", width=9)
        sx = lambda v: x1 + (v - amin) / (amax - amin) * (x2 - x1)
        d.line([(sx(low), y + 12), (sx(high), y + 12)], fill="#5B8F75", width=20)
        d.ellipse((sx(default) - 16, y - 4, sx(default) + 16, y + 28), fill="#D9534F", outline="white", width=4)
        d.text((sx(default) - 60, y - 50), f"默认 {default:g}", font=font(24, True), fill="#A8322E")
        d.text((x1 - 10, y + 42), f"{amin:g}", font=font(20), fill="#6B7280")
        d.text((x2 - 40, y + 42), f"{amax:g}", font=font(20), fill="#6B7280")

    d.text((70, 900), "植株高度和果实高度属于场景几何参数；留果量是按果穗数与每穗留果数推算的单茬计划值。", font=font(24), fill="#4B5563")
    im.save(path, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = tc_borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_hyperlink(paragraph, text, url, color="2F6B55", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")
    r_pr.append(rfonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_run_font(run, size=10.5, bold=False, color="111111", name="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")


def style_table(table, widths=None, header=True, font_size=9.2):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0 and header:
            set_repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = widths[c_idx]
            if r_idx == 0 and header:
                set_cell_shading(cell, GREEN)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE)
            else:
                set_cell_shading(cell, "FFFFFF")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx < 3 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, font_size, bold=(r_idx == 0), color="FFFFFF" if r_idx == 0 else "1F2933")


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    set_col_widths(table, widths)
    style_table(table, widths, header=True, font_size=font_size)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    set_run_font(r, 9.5, color="4B5563")
    return p


def add_figure(doc, image_path, width, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(image_path), width=width)
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    return p


def remove_paragraph_borders(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is not None:
        ppr.remove(pbdr)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    size = {1: 16, 2: 12.5, 3: 11.2}.get(level, 11)
    set_run_font(r, size, bold=True, color="000000")
    remove_paragraph_borders(p)
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, 10.5, bold=True)
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r, 10.2)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, 9, color="6B7280")


def build_document():
    height_fig = ASSET_DIR / "height_and_fruit_zone.png"
    layout_fig = ASSET_DIR / "greenhouse_layout.png"
    range_fig = ASSET_DIR / "parameter_ranges.png"
    make_height_figure(height_fig)
    make_layout_figure(layout_fig)
    make_range_figure(range_fig)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)

    title_style = styles["Title"]
    title_style.font.name = "Microsoft YaHei"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.paragraph_format.space_after = Pt(10)
    title_ppr = title_style._element.get_or_add_pPr()
    title_pbdr = title_ppr.find(qn("w:pBdr"))
    if title_pbdr is not None:
        title_ppr.remove(title_pbdr)

    for i, size in ((1, 16), (2, 12.5), (3, 11.2)):
        st = styles[f"Heading {i}"]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(12 if i == 1 else 8)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    # Footer
    footer = sec.footer
    p = footer.paragraphs[0]
    add_page_number(p)

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(34)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = p.add_run("番茄棚关键参数与采摘机器人场景调研")
    set_run_font(title_run, 25, bold=True, color="000000")
    remove_paragraph_borders(p)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(24)
    r = p2.add_run("用于番茄田生成参数调整与预设采摘点设计")
    set_run_font(r, 14, color="4B5563")

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(22)
    r = p3.add_run("调研结论")
    set_run_font(r, 12, bold=True, color=GREEN)

    add_body(doc, "本报告把设施番茄栽培资料、品种登记数据和采摘机器人研究转换为可直接用于仿真场景的参数。建议基准场景采用无限生长型大中果番茄、单干整枝、短至中等栽培周期。株距设为 0.40 m，可见植株高度设为 1.60 m，成熟果实中心高度覆盖 0.55–1.40 m，单株计划留果量取 24 个。机器人净通道先按 1.00–1.20 m 设置，再根据底盘宽度和安全余量复核。")

    summary_rows = [
        ("株距", "0.30–0.50 m", "0.40 m", "控制植株沿行方向密度"),
        ("可见植株高度", "1.50–1.70 m", "1.60 m", "不等同于长季栽培的累计茎长"),
        ("成熟果实高度", "0.55–1.40 m", "0.60–1.30 m", "按果实中心距地面计算"),
        ("单株留果量", "16–40 个", "24 个", "按 6 穗乘每穗 4 果"),
        ("果实横径", "62.2–98.1 mm", "78 mm", "大果型样本 100 果"),
        ("机器人净通道", "0.90–1.20 m", "1.10 m", "需满足底盘宽度和两侧余量"),
    ]
    add_table(doc, ["参数", "资料范围", "基准值", "使用说明"], summary_rows,
              [Inches(1.35), Inches(1.28), Inches(1.18), Inches(2.95)], font_size=9.0)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("2026 年 10 月 10 日")
    set_run_font(r, 10, color="6B7280")

    doc.add_page_break()

    # 1 Scope
    add_heading(doc, "调研范围与参数口径", 1)
    add_body(doc, "调研对象为日光温室和连栋温室中的鲜食番茄，重点面向仿真番茄田、移动底盘导航和机械臂采摘。资料中的栽培密度、果穗数量和果实高度会随品种、季节、整枝方式和是否落蔓而变化，因此报告同时给出资料范围和工程默认值。默认值用于建立可复现的基准场景，不代表所有番茄棚的唯一标准。")

    defs = [
        ("株距", "同一栽培行内相邻植株茎基部中心之间的距离。"),
        ("行距", "相邻栽培行中心之间的距离。大小行种植时分别记录小行距和作为作业通道的大行距。"),
        ("植株高度", "仿真场景中地面至冠层顶端的可见垂直高度。长季高线栽培会落蔓，累计茎长可能远大于可见高度。"),
        ("结果高度", "果实几何中心到地面的垂直距离。机械臂工作空间和相机升降范围应按该口径设计。"),
        ("单株结果数量", "本报告优先采用单茬计划留果量，即果穗数乘每穗保留果实数；它不等于某一时刻可见的成熟果数量。"),
    ]
    add_table(doc, ["术语", "统一口径"], defs, [Inches(1.35), Inches(5.35)], font_size=9.4)

    add_heading(doc, "证据等级", 2)
    add_bullet(doc, "A 级资料包括政府农业部门发布的栽培技术、地方标准和品种登记资料。")
    add_bullet(doc, "B 级资料包括具有明确试验环境和样本数据的同行评议论文。")
    add_bullet(doc, "工程建议由多项资料交叉形成，并在文中标为建议值或推算值。")

    add_heading(doc, "适用场景", 2)
    add_body(doc, "基准参数适合生成单干整枝、株型较规整的成熟番茄行。若项目改为樱桃番茄、长季高线栽培或露地番茄，应重新调整果实数量、果径和植株高度。机器人通道还必须结合底盘外廓、转弯半径、雷达盲区和机械臂安装位置校核。")

    doc.add_page_break()

    # 2 Sources table
    add_heading(doc, "资料中的关键参数", 1)
    source_rows = [
        ("淄博日光温室技术", "大行 0.80–0.90 m，小行 0.50–0.60 m；6–8 穗；每穗 3–5 果", "A", "栽培布局与留果量"),
        ("山西基质栽培规程", "行距 0.40 m，株距 0.50 m；4–5 穗，每穗 4–5 果", "A", "短周期基质栽培"),
        ("黑龙江长季模式", "畦宽 1.40 m，双行，株距 0.40 m；1800–2000 株每亩", "A", "无限生长型长季栽培"),
        ("北京品种登记资料", "部分品种株高约 1.70 m；常见 5–8 穗，每穗约 4–5 果", "A", "植株高度与留果量"),
        ("夜间采摘机器人试验", "行距 1.00 m，株距 0.30 m，株高 1.50 m；100 果横径均值 78.4 mm", "B", "机器人真实作业环境"),
        ("采摘机械臂研究", "成熟果实高度 0.55–1.20 m；株距 0.30 m；作业空间高度 0.65 m", "B", "机械臂工作空间"),
        ("番茄采摘底盘研究", "株高 0.30–1.60 m，株距 0.26–0.82 m，垄间距约 0.80 m", "B", "底盘和导航边界"),
        ("机械化布局试验", "密植与低密度方案包含 0.90–1.20 m 主作业间距", "B", "机械通行布局"),
    ]
    add_table(doc, ["资料", "报告采用的数据", "等级", "主要用途"], source_rows,
              [Inches(1.55), Inches(3.45), Inches(0.55), Inches(1.25)], font_size=8.65)

    add_body(doc, "资料差异主要来自品种和栽培周期。短周期栽培常在 4–6 穗后打顶，长季栽培会持续生长并采用落蔓管理。为了同时服务导航和采摘，基准场景采用 1.60 m 可见高度，同时保留 0.55–1.40 m 的果实工作区。")

    add_figure(doc, range_fig, Inches(6.75), "株距 植株高度 成熟果高度 单株留果量和果实横径的文献范围与推荐默认值")
    add_caption(doc, "图 1  核心参数资料范围与建议默认值")

    doc.add_page_break()

    # 3 Core parameters
    add_heading(doc, "番茄植株与果实参数", 1)
    add_heading(doc, "株距与植株高度", 2)
    add_body(doc, "政府技术规程和机器人试验中的株距主要落在 0.30–0.50 m。基质栽培地方标准使用 0.50 m，采摘机器人试验使用 0.30 m，长季栽培模式常见 0.40 m。基准值取 0.40 m，可以兼顾植株密度和视觉遮挡，并便于按固定步长生成植株。")
    add_body(doc, "成熟植株的可见高度可按 1.50–1.70 m 建模。江苏普罗旺斯番茄采摘试验采用 1.50 m，品种登记资料给出约 1.70 m 的株高，底盘研究覆盖至 1.60 m。基准值取 1.60 m，随机扰动建议设为正负 0.15 m。若后续采用高线落蔓场景，应单独建模茎蔓路径，不宜仅拉高整株模型。")

    add_heading(doc, "结果高度", 2)
    add_body(doc, "机械臂研究把成熟果实高度设为 0.55–1.20 m，另一类温室机器人研究将番茄主要分布区概括为约 0.60–1.40 m。综合后建议场景允许成熟果实中心出现在 0.55–1.40 m，并把 0.60–1.30 m 作为常规生成范围。若采用概率分布，可用截断正态分布，均值设为 0.95 m，标准差设为 0.22 m，截断在 0.55 m 和 1.40 m。该分布属于工程建议，应在实棚测量后更新。")

    add_figure(doc, height_fig, Inches(6.65), "番茄植株高度与果实中心结果高度的测量口径示意图")
    add_caption(doc, "图 2  植株高度与果实中心高度的统一口径  数据范围见参考资料 4 至 7")

    doc.add_page_break()

    # 4 Fruit count
    add_heading(doc, "单株结果数量与成熟比例", 1)
    add_body(doc, "单株结果数量必须区分计划留果量、瞬时挂果量和成熟可采量。农业技术资料通常给出果穗数及每穗留果数。短周期基质栽培为 4–5 穗、每穗 4–5 果，对应 16–25 个；温室管理资料给出 6–8 穗、每穗 3–5 果，对应约 18–40 个；部分登记品种采用 6–8 穗、每穗 4–5 果，对应 24–40 个。")

    count_rows = [
        ("单茬计划留果量", "16–40 个", "24 个", "6 穗乘每穗 4 果，适合作为植株总果实上限"),
        ("场景瞬时挂果量", "12–28 个", "20 个", "考虑下部已采收和上部未坐果，属于工程推算"),
        ("成熟可采果", "3–10 个", "5 个", "按瞬时挂果量的约 25% 设置成熟状态，需按任务调整"),
        ("单穗果实数", "3–5 个", "4 个", "大中果型番茄常见疏果目标"),
        ("果穗数量", "4–8 穗", "6 穗", "短周期与中等周期场景的折中值"),
    ]
    add_table(doc, ["数量口径", "建议范围", "默认值", "说明"], count_rows,
              [Inches(1.4), Inches(1.1), Inches(0.9), Inches(3.35)], font_size=9.1)

    add_heading(doc, "场景生成方法", 2)
    add_bullet(doc, "先生成 6 个果穗，果穗中心沿高度方向分布在 0.60–1.30 m。")
    add_bullet(doc, "每穗生成 3–5 个果实，平均为 4 个；单株初始总数约 24 个。")
    add_bullet(doc, "随机选择约 20% 至 30% 的果实作为成熟可采状态，其余分为转色和青熟状态。")
    add_bullet(doc, "为增加自然度，可让同一果穗的果实高度差控制在 0.05–0.18 m，并设置少量遮挡。")

    add_heading(doc, "果实尺寸", 2)
    add_body(doc, "普罗旺斯番茄试验随机测量 100 个果实，横径为 62.2–98.1 mm，均值 78.4 mm；纵径为 52.5–72.0 mm，均值 61.0 mm；质量为 181.9–285.4 g，均值 209.1 g。大果型基准场景可将横径设为 78 mm、纵径设为 61 mm，并分别加入约 10% 的随机扰动。樱桃番茄需要另建尺寸和数量配置。")

    fruit_rows = [
        ("横径", "62.2–98.1 mm", "78.4 mm", "碰撞体与抓取开口"),
        ("纵径", "52.5–72.0 mm", "61.0 mm", "果实模型纵向尺度"),
        ("质量", "181.9–285.4 g", "209.1 g", "末端负载和收集筐估算"),
    ]
    add_table(doc, ["指标", "实测范围", "实测均值", "仿真用途"], fruit_rows,
              [Inches(1.05), Inches(1.55), Inches(1.35), Inches(2.8)], font_size=9.25)

    doc.add_page_break()

    # 5 Layout
    add_heading(doc, "番茄棚布局与导航空间", 1)
    add_body(doc, "生产资料常用大小行或双行栽培。淄博技术要点给出大行 0.80–0.90 m、小行 0.50–0.60 m；采摘机器人试验使用 1.00 m 行距；机械化布局研究把主作业间距扩展到 0.90–1.20 m。导航场景应把作物行中心距和底盘可通行净宽分别建模，因为冠层会侵占名义行距。")

    layout_rows = [
        ("小行距", "0.40–0.60 m", "0.50 m", "双行种植内部间距"),
        ("机器人净通道", "0.90–1.20 m", "1.10 m", "冠层边界之间的实际空隙"),
        ("株距", "0.30–0.50 m", "0.40 m", "沿行方向"),
        ("垄高", "0.15–0.35 m", "0.20 m", "技术资料和基质槽结构的折中值"),
        ("冠层半宽", "0.25–0.40 m", "0.32 m", "建议参数，决定通道实际宽度"),
    ]
    add_table(doc, ["几何参数", "建议范围", "默认值", "定义"], layout_rows,
              [Inches(1.3), Inches(1.2), Inches(1.0), Inches(3.25)], font_size=9.2)

    add_figure(doc, layout_fig, Inches(6.75), "番茄双行种植布局 机器人净通道 株距与预设采摘点示意图")
    add_caption(doc, "图 3  双行种植与机器人作业通道的基准布局")

    add_heading(doc, "通道校核", 2)
    add_body(doc, "净通道宽度应不小于底盘宽度加两侧安全余量。建议每侧保留 0.15–0.25 m；若底盘宽 0.65 m，净通道宜为 0.95–1.15 m。转弯区还需要额外长度，不能只用直行通道宽度判断是否可达。地图中的障碍物边界应按冠层外缘建立，而不是按茎秆中心建立。")

    doc.add_page_break()

    # 6 Robot research
    add_heading(doc, "农业采摘机器人调研结论", 1)
    add_body(doc, "现有番茄采摘机器人通常由移动底盘、环境感知、机械臂、末端执行器和果实收集模块组成。真实温室试验表明，果实高度跨度、叶片遮挡和狭窄通道会同时影响识别、停车和采摘。因此，场景生成不能只放置果实模型，还应包含冠层体积、垄体、通道边界和不同成熟状态。")

    robot_rows = [
        ("导航", "激光 SLAM 或视觉激光融合；地图定位；全局和局部路径规划", "通道边界、行尾转弯区、临时障碍物"),
        ("果实感知", "RGB 或 RGB D 相机识别成熟度并估计三维位置", "遮挡比例、光照变化、果实尺寸和颜色"),
        ("机械臂", "固定或升降安装，覆盖约 0.55–1.40 m 果实高度", "工作空间、避碰、伸展距离和安装高度"),
        ("末端执行器", "夹持、剪切或旋拧，需适应约 62–98 mm 大果横径", "抓取开口、果柄方向、邻果碰撞"),
        ("任务规划", "底盘停车后分层观察和采摘，再移动到下一作业点", "采摘点间隔、面向行方向和覆盖植株数量"),
    ]
    add_table(doc, ["模块", "典型做法", "场景需要提供的信息"], robot_rows,
              [Inches(1.05), Inches(2.75), Inches(2.95)], font_size=9.0)

    add_heading(doc, "预设采摘点建议", 2)
    add_body(doc, "以 0.40 m 株距为基准，可先每两株设置一个停车点，即点间距 0.80 m。停车点位于通道中心线，朝向番茄行，左右两侧分别建立作业点或通过机械臂基座旋转覆盖。文献中的机械臂最大工作距离可达约 0.875 m，但实际有效距离会受到安装高度、叶片和碰撞约束，因此停车点覆盖范围应通过仿真逐点验证。")
    add_bullet(doc, "位置误差验收建议为平面距离不超过 0.10 m，航向误差不超过 5 度。")
    add_bullet(doc, "每个采摘点记录位置、朝向、服务的植株编号和左右作业侧。")
    add_bullet(doc, "在行首、行尾和狭窄位置增加禁停区，避免机械臂展开时侵入障碍物。")
    add_bullet(doc, "默认地图完成后，用全部预设点做批量可达性检查，并记录失败点及原因。")

    add_heading(doc, "机器人调研对当前开发的影响", 2)
    add_body(doc, "当前的建图和底盘到目标点功能可以在现有场景中完成。参数调研完成后，应按本报告生成默认番茄田，再重新建图并配置采摘点。采摘点是底盘停车位，不是果实坐标；果实坐标由感知或场景真值提供。")

    doc.add_page_break()

    # 7 Implementation config
    add_heading(doc, "场景生成参数建议", 1)
    add_body(doc, "下表可直接转换为配置文件。随机参数建议使用固定种子生成基准地图，同时保留至少两个变化场景测试导航和采摘算法。")

    config_rows = [
        ("plant_spacing_m", "0.40", "0.30–0.50", "正态或均匀扰动，最小间距不得低于碰撞体直径"),
        ("plant_height_m", "1.60", "1.45–1.75", "截断正态分布，标准差约 0.10"),
        ("canopy_half_width_m", "0.32", "0.25–0.40", "用于障碍物边界和激光可见轮廓"),
        ("fruit_height_min_m", "0.55", "固定", "成熟果实中心下限"),
        ("fruit_height_max_m", "1.40", "固定", "成熟果实中心上限"),
        ("fruit_height_mean_m", "0.95", "0.85–1.05", "截断正态分布中心，工程建议"),
        ("truss_count", "6", "4–8", "整数离散分布"),
        ("fruit_per_truss", "4", "3–5", "整数离散分布"),
        ("fruit_count_per_plant", "24", "16–40", "计划留果量"),
        ("ripe_ratio", "0.25", "0.15–0.40", "任务难度参数，非固定农艺指标"),
        ("fruit_diameter_m", "0.078", "0.062–0.098", "大果型横径"),
        ("clear_aisle_width_m", "1.10", "0.90–1.20", "按实际底盘宽度复核"),
        ("harvest_pose_step_m", "0.80", "0.40–1.20", "默认每两株一个停车点"),
    ]
    add_table(doc, ["参数名", "默认值", "允许范围", "生成规则"], config_rows,
              [Inches(1.9), Inches(0.9), Inches(1.15), Inches(2.8)], font_size=8.45)

    add_heading(doc, "建议测试场景", 2)
    test_rows = [
        ("基准场景", "株距 0.40 m，通道 1.10 m，株高 1.60 m，单株 24 果", "建立默认地图与预设采摘点"),
        ("密植场景", "株距 0.30 m，通道 0.90 m，冠层半宽 0.38 m", "验证狭窄通道和遮挡"),
        ("稀植场景", "株距 0.50 m，通道 1.20 m，株高波动增大", "验证定位和点位泛化"),
        ("高负载场景", "8 穗，每穗 5 果，成熟比例 0.40", "验证感知和任务调度"),
    ]
    add_table(doc, ["场景", "关键参数", "用途"], test_rows,
              [Inches(1.15), Inches(3.65), Inches(1.95)], font_size=9.1)

    add_heading(doc, "实棚复测清单", 2)
    add_body(doc, "文献参数足以建立初版默认场景，但最终模型应使用目标温室的实测数据校正。建议随机选择不少于 30 株，测量株距、冠幅和株高；每株记录全部果实中心高度、成熟度和可见性；同时测量通道最窄处、行尾空间以及底盘外廓。实测后优先更新结果高度分布、冠层宽度和采摘点间距。")

    doc.add_page_break()

    # 8 References
    add_heading(doc, "参考资料", 1)
    refs = [
        ("淄博市数字农业农村发展中心  日光温室番茄生产管理技术要点  2022", "https://ny.zibo.gov.cn/art/2022/8/24/art_980_2327495.html"),
        ("秦安县人民政府  日光温室番茄栽培技术要点", "https://www.qinan.gov.cn/info/10632/194301.htm"),
        ("山西省地方标准 DB14 T 2337 2021  设施番茄简易基质栽培技术规程", "https://www.horticulture.cn/p-15938.html"),
        ("黑龙江省农业农村厅  2025 年农作物栽培技术模式汇编  大棚番茄长季栽培技术模式", "https://nynct.hlj.gov.cn/nynct/c115423/202502/31810960/files/2025%E5%B9%B4%E9%BB%91%E9%BE%99%E6%B1%9F%E7%9C%81%E5%86%9C%E4%BD%9C%E7%89%A9%E6%A0%BD%E5%9F%B9%E6%8A%80%E6%9C%AF%E6%A8%A1%E5%BC%8F%E6%B1%87%E7%BC%96.pdf"),
        ("北京市农业农村局相关品种登记资料汇编  2022", "https://www.beijing.gov.cn/zhengce/zhengcefagui/qtwj/202204/W020220417857831999352.pdf"),
        ("Liu L et al  Design and Experiment of Nighttime Greenhouse Tomato Harvesting Robot  2024", "https://doi.org/10.5614/j.eng.technol.sci.2024.56.3.3"),
        ("Structure Optimization and Path Planning of Tomato Picking Manipulator  2016", "https://doi.org/10.1109/ISCID.2016.196"),
        ("Movement Mechanism of the Greenhouse Tomato Picking Robot Chassis  Agriculture  2023", "https://doi.org/10.3390/agriculture13030532"),
        ("Autonomous navigation system for greenhouse tomato picking robots based on laser SLAM  2024", "https://doi.org/10.1016/j.jksues.2024.102104"),
        ("Evaluating the canopy light environment photosynthesis and fruit performance under mechanized planting layouts  2025", "https://doi.org/10.1016/j.hpj.2025.03.010"),
        ("Autonomous Tomato Harvesting System Integrating AI Controlled Robotics in Greenhouses  2026", "https://doi.org/10.3390/agriculture16080847"),
    ]
    for idx, (title, url) in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.0)
        p.paragraph_format.first_line_indent = Inches(-0.0)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.18
        r = p.add_run(f"{idx}. ")
        set_run_font(r, 9.6, bold=True, color=GREEN)
        add_hyperlink(p, title, url)

    add_heading(doc, "使用说明", 2)
    add_body(doc, "报告中的原始范围来自上述资料；默认值和概率分布是为仿真与机器人测试形成的工程建议。实棚参数一旦可用，应保留本报告作为初始依据，并用实测统计量替换相应默认值。")

    # Core properties
    props = doc.core_properties
    props.title = "番茄棚关键参数与采摘机器人场景调研"
    props.subject = "设施番茄场景生成与采摘机器人参数"
    props.author = "调研整理"
    props.keywords = "番茄棚, 株距, 植株高度, 结果高度, 采摘机器人, 仿真参数"

    doc.save(OUT_PATH)
    print(str(OUT_PATH))


if __name__ == "__main__":
    build_document()
